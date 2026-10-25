import hashlib
import importlib.util
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("files", ROOT / "export-runtime" / "flowly_files.py")
files = importlib.util.module_from_spec(spec)
spec.loader.exec_module(files)

class FilesTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="flowly-files-")
        self.root = Path(self.temp.name)
        (self.root / "inputs").mkdir()
    def tearDown(self):
        self.temp.cleanup()
    def call(self, operation, config, source, node="node"):
        return files.execute({"operation": operation, "root": str(self.root), "stateDirectory": str(self.root / ".state"), "runId": "run-safe", "nodeId": node, "config": config, "input": source})
    def input(self, name):
        return self.call("input", {"path": "file", "maxBytes": 10000000}, {"file": name})
    def extract(self, data, mode="text", language="eng"):
        return self.call("extract", {"filePath": "file", "mode": mode, "maxPages": 100, "language": language}, data)
    def test_path_hash_size_and_changed_input(self):
        (self.root / "inputs" / "hello.txt").write_text("안녕하세요", encoding="utf-8")
        value = self.input("hello.txt")
        self.assertEqual(value["file"]["sha256"], hashlib.sha256("안녕하세요".encode()).hexdigest())
        self.assertEqual(self.extract(value)["text"], "안녕하세요")
        for path in ("../outside.txt", "/tmp/secret.txt", "C:\\secret.txt", "\\\\server\\secret.txt", ".env"):
            with self.assertRaises(files.FileFault): self.input(path)
        with self.assertRaises(files.FileFault): self.call("input", {"path": "file", "maxBytes": 1}, {"file": "hello.txt"})
        (self.root / "inputs" / "hello.txt").write_text("changed")
        with self.assertRaises(files.FileFault): self.extract(value)
    def test_csv_and_xlsx_roundtrip_formula_and_duplicate_write(self):
        if not importlib.util.find_spec("openpyxl"): self.skipTest("optional requirements-files.txt is not installed")
        data = [{"이름": "민수", "값": "=HYPERLINK(\"bad\")"}, {"이름": "영희", "값": 42}]
        for format in ("csv", "xlsx"):
            config = {"format": format, "filename": "결과." + format, "dataPath": "data"}
            first = self.call("output", config, {"data": data}, format)
            target = self.root / first["file"]["path"]
            stamp = target.stat().st_mtime_ns
            self.assertEqual(self.call("output", config, {"data": data}, format), first)
            self.assertEqual(stamp, target.stat().st_mtime_ns)
            extracted = self.extract(first, "table")
            self.assertEqual(extracted["tables"][0]["rows"][1][1], "'=HYPERLINK(\"bad\")")
            with self.assertRaises(files.FileFault): self.call("output", config, {"data": []}, format)
    def test_word_and_pdf_korean_output_extract(self):
        if not all(importlib.util.find_spec(p) for p in ("docx", "reportlab", "pypdf")): self.skipTest("optional document packages are not installed")
        text = "안녕하세요 Flowly\n실제 파일 출력과 문서 추출 확인"
        for format in ("docx", "pdf"):
            value = self.call("output", {"format": format, "filename": "result." + format, "dataPath": "text"}, {"text": text}, format)
            result = self.extract(value)
            self.assertIn("안녕하세요", result["text"])
            self.assertIn("Flowly", result["text"])
            if format == "docx": self.assertIsNone(result["pages"][0]["page"])
    def test_real_document_tables(self):
        if not all(importlib.util.find_spec(p) for p in ("docx", "reportlab", "pdfplumber")): self.skipTest("optional table packages are not installed")
        from docx import Document
        from reportlab.pdfgen.canvas import Canvas
        doc = Document()
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "name", "value"
        t.cell(1, 0).text, t.cell(1, 1).text = "sample", "42"
        doc.save(self.root / "inputs" / "table.docx")
        self.assertEqual(self.extract(self.input("table.docx"), "table")["tables"][0]["rows"][1], ["sample", "42"])
        canvas = Canvas(str(self.root / "inputs" / "table.pdf"))
        for x in (50, 150, 250): canvas.line(x, 650, x, 710)
        for y in (650, 680, 710): canvas.line(50, y, 250, y)
        canvas.drawString(60, 690, "name"); canvas.drawString(160, 690, "value")
        canvas.drawString(60, 660, "sample"); canvas.drawString(160, 660, "42")
        canvas.save()
        rows = self.extract(self.input("table.pdf"), "table")["tables"][0]["rows"]
        self.assertEqual(rows, [["name", "value"], ["sample", "42"]])
    def test_real_ocr_and_missing_language(self):
        if not importlib.util.find_spec("PIL") or not (importlib.util.find_spec("tesserocr") or shutil.which("tesseract")) or not (ROOT / "tessdata" / "kor.traineddata").is_file() or not Path("C:/Windows/Fonts/malgun.ttf").is_file(): self.skipTest("Korean OCR engine, language data and font are required")
        from PIL import Image, ImageDraw, ImageFont
        shutil.copytree(ROOT / "tessdata", self.root / "tessdata")
        image = Image.new("RGB", (1000, 180), "white")
        font = ImageFont.truetype("C:/Windows/Fonts/malgun.ttf", 45)
        ImageDraw.Draw(image).text((25, 40), "안녕하세요 Flowly 12345", font=font, fill="black")
        image.save(self.root / "inputs" / "ocr.png")
        result = self.extract(self.input("ocr.png"), "ocr", "kor+eng")
        self.assertIn("Flowly", result["text"])
        self.assertIn("12345", result["text"])
        self.assertIn("안녕하세요", result["text"].replace(" ", ""))
        with self.assertRaises(files.FileFault) as failure: self.extract(self.input("ocr.png"), "ocr", "missing")
        self.assertEqual(failure.exception.kind, "unconfigured")
    def test_limits_and_output_collision(self):
        (self.root / "inputs" / "bad.pdf").write_bytes(b"not a pdf")
        with self.assertRaises(Exception): self.extract(self.input("bad.pdf"))
        with self.assertRaises(files.FileFault): self.call("output", {"format": "csv", "filename": "../escape.csv", "dataPath": "data"}, {"data": []})
        with self.assertRaises(files.FileFault): self.call("output", {"format": "csv", "filename": "big.csv", "dataPath": "data"}, {"data": [[1] * 10001]})
        cfg = {"format": "csv", "filename": "result.csv", "dataPath": "data"}
        value = self.call("output", cfg, {"data": [["safe"]]})
        (self.root / value["file"]["path"]).write_bytes(b"changed")
        with self.assertRaises(files.FileFault): self.call("output", cfg, {"data": [["safe"]]})
    def test_python_311_windows_reparse_points_are_rejected(self):
        linked = SimpleNamespace(parents=[], is_symlink=lambda: False, lstat=lambda: SimpleNamespace(st_file_attributes=0x400))
        with self.assertRaises(files.FileFault): files.no_links(linked)
    def test_ocr_check_requires_python_wrapper_for_binary_fallback(self):
        (self.root / "tessdata").mkdir()
        (self.root / "tessdata" / "eng.traineddata").write_bytes(b"fixture only")
        with patch.object(files.importlib.util, "find_spec", side_effect=lambda name: None if name in ("tesserocr", "pytesseract") else object()), patch.object(files.shutil, "which", return_value="tesseract"):
            checked=files.check({"nodes":[{"type":"document_extract","config":{"mode":"ocr","language":"eng"}}]},self.root)
        self.assertEqual(checked["status"],"unconfigured")
        self.assertIn("pytesseract",checked["missing"])

if __name__ == "__main__": unittest.main()
